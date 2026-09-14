import csv
import importlib.util
import io
import pkgutil
import re

if not hasattr(pkgutil, "get_loader"):
    def _compat_get_loader(module_name):
        try:
            spec = importlib.util.find_spec(module_name)
            if spec is None:
                return None
            return spec.loader
        except (ImportError, AttributeError, ValueError):
            return None

    pkgutil.get_loader = _compat_get_loader

from flask import Flask, redirect, render_template, request, session, url_for

from logic import compute_stats, get_alerts, get_status
from mock_data import CUSTOMERS, NETWORK_POSTS, SALES_ORDERS, STOCK_LEVELS, TEAM, TRANSFERS

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

app = Flask(__name__)
app.secret_key = "jodoo-dev-secret"
    from auth.routes import auth_bp
    # use dashboard1 as requested
    from dashboard1.routes import dashboard_bp as dashboard_bp

COLUMN_ALIASES = {
    "product": ["product", "product name", "item", "item name", "name"],
    "sku": ["sku", "sku id", "product code", "code", "item code"],
    "location": ["location", "warehouse", "warehouse id", "warehouse name", "store", "site"],
    "qty": ["qty", "quantity", "stock", "current stock", "on hand", "available stock"],
    "min_qty": ["min quantity", "minimum stock", "reorder point", "reorder level", "threshold", "safety stock", "minimum qty"],
}

SALES_COLUMN_ALIASES = {
    "ref": ["transaction id", "order id", "receipt number", "ref", "invoice id", "invoice number"],
    "product": ["item", "item name", "product", "product name"],
    "quantity": ["quantity", "qty", "units", "count"],
    "unit_price": ["price per unit", "unit price", "price"],
    "amount": ["total spent", "total", "amount", "revenue", "sales value"],
    "date": ["transaction date", "date", "order date", "sale date"],
    "customer": ["customer", "customer name"],
    "payment_method": ["payment method", "payment type"],
    "location": ["location", "store", "warehouse"],
}

FIELD_LABELS = {
    "product": "product",
    "sku": "SKU",
    "location": "location",
    "qty": "quantity",
    "min_qty": "minimum quantity",
}

SALES_FIELD_LABELS = {
    "ref": "transaction reference",
    "product": "item/product",
    "date": "transaction date",
    "amount": "amount",
    "quantity": "quantity",
    "unit_price": "unit price",
}


def with_status(stock):
    return [{**item, "status": get_status(item)} for item in stock]


def normalize_header(value):
    if value is None:
        return ""
    text = str(value).lower()
    text = text.replace("_", " ").replace("-", " ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def match_header_score(header_text, alias_text):
    header = normalize_header(header_text)
    alias = normalize_header(alias_text)
    if not header or not alias:
        return 0
    if header == alias:
        return 100
    if alias in header or header in alias:
        return 80

    header_tokens = set(header.split())
    alias_tokens = set(alias.split())
    if alias_tokens and header_tokens and alias_tokens.issubset(header_tokens):
        return 60
    if alias_tokens and header_tokens and header_tokens.issubset(alias_tokens):
        return 60
    return 0


def find_best_header_for_concept(detected_headers, concept, alias_map, used_headers=None):
    used_headers = used_headers or set()
    best_header = None
    best_rank = None

    for index, header in enumerate(detected_headers):
        if header in used_headers:
            continue
        normalized_header = normalize_header(header)
        if not normalized_header:
            continue

        for alias_index, alias in enumerate(alias_map.get(concept, [])):
            score = match_header_score(normalized_header, alias)
            if score <= 0:
                continue
            rank = (score, -alias_index, -index)
            if best_rank is None or rank > best_rank:
                best_rank = rank
                best_header = header

    return best_header


def get_active_stock():
    uploaded_stock = session.get("uploaded_stock")
    if uploaded_stock:
        return uploaded_stock
    return STOCK_LEVELS


def get_active_sales():
    uploaded_sales = session.get("uploaded_sales")
    if uploaded_sales:
        return uploaded_sales
    return SALES_ORDERS


def count_header_matches(detected_headers, alias_map):
    matches = 0
    for concept, aliases in alias_map.items():
        for header in detected_headers:
            if any(match_header_score(header, alias) > 0 for alias in aliases):
                matches += 1
                break
    return matches


def detect_upload_type(detected_headers):
    stock_score = 0
    for concept in ["product", "sku", "location", "qty", "min_qty"]:
        if any(match_header_score(header, alias) > 0 for header in detected_headers for alias in COLUMN_ALIASES.get(concept, [])):
            stock_score += 1

    sales_score = 0
    for concept in ["product", "date", "amount", "quantity", "unit_price"]:
        if any(match_header_score(header, alias) > 0 for header in detected_headers for alias in SALES_COLUMN_ALIASES.get(concept, [])):
            sales_score += 1

    if stock_score == 0 and sales_score == 0:
        return "unknown"
    if sales_score > stock_score:
        return "sales"
    if stock_score > sales_score:
        return "stock"
    return "unknown"


def parse_stock_rows_from_file(uploaded_file):
    if uploaded_file is None or not getattr(uploaded_file, "filename", ""):
        raise ValueError("Please upload a .csv or .xlsx file")

    extension = uploaded_file.filename.rsplit(".", 1)[-1].lower() if "." in uploaded_file.filename else ""

    if extension == "csv":
        uploaded_file.stream.seek(0)
        text_stream = io.TextIOWrapper(uploaded_file.stream, encoding="utf-8-sig", newline="")
        reader = csv.DictReader(text_stream)
        raw_rows = list(reader)
        detected_headers = list(reader.fieldnames or [])
        return raw_rows, detected_headers

    if extension == "xlsx":
        if load_workbook is None:
            raise ValueError("Excel upload support is unavailable because openpyxl is not installed.")
        uploaded_file.stream.seek(0)
        workbook = load_workbook(io.BytesIO(uploaded_file.read()), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("The uploaded Excel file is empty.")
        headers = [str(cell).strip() for cell in rows[0] if cell is not None]
        raw_rows = []
        for row in rows[1:]:
            if row is None or not any((value is not None and str(value).strip() != "") for value in row):
                continue
            raw_rows.append({headers[index]: row[index] for index in range(min(len(headers), len(row)))})
        return raw_rows, headers

    raise ValueError("Please upload a .csv or .xlsx file")


def map_stock_columns(raw_rows, detected_headers):
    matched_columns = {}
    used_headers = set()
    missing_columns = []

    for canonical_name in ["product", "sku", "location", "qty", "min_qty"]:
        match = find_best_header_for_concept(detected_headers, canonical_name, COLUMN_ALIASES, used_headers)
        if match is None:
            missing_columns.append(f"a {FIELD_LABELS[canonical_name]} column")
        else:
            matched_columns[canonical_name] = match
            used_headers.add(match)

    if missing_columns:
        missing_text = "; ".join(missing_columns)
        raise ValueError(f"Could not find {missing_text}. Detected columns: {', '.join(detected_headers) if detected_headers else 'none'}")

    return matched_columns


def build_stock_records(raw_rows, matched_columns):
    parsed_rows = []
    for row_number, row in enumerate(raw_rows, start=2):
        if row is None or not any((value is not None and str(value).strip() != "") for value in row.values()):
            continue

        mapped_row = {}
        for canonical_name, actual_header in matched_columns.items():
            value = row.get(actual_header, "")
            mapped_row[canonical_name] = value

        try:
            qty = int(str(mapped_row["qty"]).strip())
            min_qty = int(str(mapped_row["min_qty"]).strip())
        except (TypeError, ValueError):
            raise ValueError(f"Row {row_number} has a non-numeric quantity or minimum quantity. Use whole numbers only.")

        product = str(mapped_row.get("product", "") or "").strip()
        sku = str(mapped_row.get("sku", "") or "").strip()
        location = str(mapped_row.get("location", "") or "").strip()

        if not product or not sku or not location:
            raise ValueError(f"Row {row_number} is missing a required product, SKU, or location value.")

        parsed_rows.append({
            "product": product,
            "sku": sku,
            "location": location,
            "qty": qty,
            "min_qty": min_qty,
        })

    if not parsed_rows:
        raise ValueError("The uploaded file contains no valid stock rows.")

    return parsed_rows


def map_sales_columns(detected_headers):
    matched_columns = {}
    used_headers = set()

    for canonical_name in ["ref", "product", "quantity", "unit_price", "amount", "date", "customer", "payment_method", "location"]:
        match = find_best_header_for_concept(detected_headers, canonical_name, SALES_COLUMN_ALIASES, used_headers)
        if match is not None:
            matched_columns[canonical_name] = match
            used_headers.add(match)

    return matched_columns


def build_sales_records(raw_rows, matched_columns):
    parsed_rows = []
    generated_ref = 1

    for row_number, row in enumerate(raw_rows, start=2):
        if row is None or not any((value is not None and str(value).strip() != "") for value in row.values()):
            continue

        product = str(row.get(matched_columns.get("product", ""), "") or "").strip()
        date_value = str(row.get(matched_columns.get("date", ""), "") or "").strip()
        customer = str(row.get(matched_columns.get("customer", ""), "") or "").strip() or "Unknown"
        location = str(row.get(matched_columns.get("location", ""), "") or "").strip() or ""
        payment_method = str(row.get(matched_columns.get("payment_method", ""), "") or "").strip() or ""

        if not product or not date_value:
            raise ValueError(f"Row {row_number} is missing a required product or date value for a sales/transaction file.")

        quantity_value = row.get(matched_columns.get("quantity", ""), "")
        unit_price_value = row.get(matched_columns.get("unit_price", ""), "")
        amount_value = row.get(matched_columns.get("amount", ""), "")

        try:
            quantity = float(str(quantity_value).strip()) if quantity_value not in (None, "") else None
        except (TypeError, ValueError):
            quantity = None

        try:
            unit_price = float(str(unit_price_value).strip()) if unit_price_value not in (None, "") else None
        except (TypeError, ValueError):
            unit_price = None

        try:
            amount = float(str(amount_value).strip()) if amount_value not in (None, "") else None
        except (TypeError, ValueError):
            amount = None

        if amount is None and quantity is not None and unit_price is not None:
            amount = quantity * unit_price

        if amount is None:
            raise ValueError(f"Row {row_number} is missing a valid amount value, or a valid quantity + unit price pair, for the sales file.")

        ref_value = str(row.get(matched_columns.get("ref", ""), "") or "").strip()
        if not ref_value:
            ref_value = f"TXN-{generated_ref:04d}"
            generated_ref += 1
        else:
            generated_ref += 1

        parsed_rows.append({
            "ref": ref_value,
            "customer": customer,
            "amount": round(float(amount), 2),
            "status": "Completed",
            "date": date_value,
            "items": product,
            "payment_method": payment_method,
            "location": location,
        })

    if not parsed_rows:
        raise ValueError("The uploaded file contains no valid sales/transaction rows.")

    return parsed_rows


@app.get("/")
def landing_page():
    # Public landing page for the product marketing site.
    return render_template("landing.html")


@app.route("/login", methods=["GET", "POST"])
def login_page():
    # Login form: fake session auth to keep the prototype working without real auth.
    if request.method == "POST":
        session["logged_in"] = True
        return redirect(url_for("onboarding_page"))
    return render_template("login.html")


@app.get("/onboarding")
def onboarding_page():
    # If the user already completed onboarding in this session, skip to the upload step.
    if session.get("onboarding"):
        return redirect(url_for("upload_data_page"))
    return render_template("questionnaire.html")


@app.post("/onboarding")
def onboarding_submit():
    answers = {
        "business_type": request.form.get("business_type", ""),
        "operations": request.form.getlist("operations"),
        "contact_channels": request.form.getlist("contact_channels"),
        "team_size": request.form.get("team_size", ""),
        "finance_method": request.form.get("finance_method", ""),
        "goals": request.form.getlist("goals"),
        "workspace_pref": request.form.get("workspace_pref", ""),
    }
    session["onboarding"] = answers
    return redirect(url_for("upload_data_page"))


@app.get("/upload-data")
def upload_data_page():
    return render_template("upload_data.html")


@app.post("/upload-data")
def upload_data_submit():
    uploaded_file = request.files.get("stock_csv")

    # Skip without a file means stay on the original sample-data fallback.
    if not uploaded_file or uploaded_file.filename == "":
        session.pop("uploaded_stock", None)
        session.pop("uploaded_sales", None)
        return redirect(url_for("dashboard_page"))

    try:
        raw_rows, detected_headers = parse_stock_rows_from_file(uploaded_file)

        if not detected_headers:
            raise ValueError("The uploaded file is empty or missing a header row.")

        upload_type = detect_upload_type(detected_headers)

        if upload_type == "stock":
            matched_columns = map_stock_columns(raw_rows, detected_headers)
            parsed_rows = build_stock_records(raw_rows, matched_columns)
            session.pop("uploaded_sales", None)
            session["uploaded_stock"] = parsed_rows
            return redirect(url_for("dashboard_page"))

        if upload_type == "sales":
            matched_columns = map_sales_columns(detected_headers)
            parsed_rows = build_sales_records(raw_rows, matched_columns)
            session.pop("uploaded_stock", None)
            session["uploaded_sales"] = [
                {
                    "ref": row["ref"],
                    "customer": row["customer"],
                    "amount": row["amount"],
                    "status": row["status"],
                    "date": row["date"],
                    "items": row["items"],
                }
                for row in parsed_rows
            ]
            return redirect(url_for("dashboard_page"))

        actual_columns = ", ".join(detected_headers) if detected_headers else "none"
        raise ValueError(
            f"We could not match this file to either a stock/inventory format or a sales/transaction format. Detected columns: {actual_columns}. "
            "Accepted stock format: product, sku, location, qty, min_qty. Accepted sales format: Transaction ID, Item, Quantity, Price Per Unit, Total Spent, Payment Method, Location, Transaction Date."
        )
    except (csv.Error, ValueError, TypeError) as exc:
        return render_template("upload_data.html", error=str(exc))


@app.get("/app/dashboard")
def dashboard_page():
    # Dashboard page uses uploaded stock data when present; otherwise it falls back to mock_data.py.
    has_uploaded_stock = bool(session.get("uploaded_stock"))
    has_uploaded_sales = bool(session.get("uploaded_sales"))

    if not has_uploaded_stock and not has_uploaded_sales:
        banner_text = "Showing sample data — upload your own on Settings"
    elif has_uploaded_stock and has_uploaded_sales:
        banner_text = "Showing your uploaded data."
    elif has_uploaded_stock:
        banner_text = "Showing your uploaded stock data. Sales are still sample data — upload them on Settings."
    else:
        banner_text = "Showing your uploaded sales data. Inventory is still sample data — upload it on Settings."

    stock = with_status(get_active_stock())
    alerts = get_alerts(stock)
    stats = compute_stats(stock, TRANSFERS)

    uploaded_sales = session.get("uploaded_sales")
    if uploaded_sales:
        revenue = sum(float(order.get("amount", 0) or 0) for order in uploaded_sales)
        order_count = len(uploaded_sales)
    else:
        revenue = 1875000
        order_count = 1246

    kpis = {
        "revenue": revenue,
        "orders": order_count,
        "profit": 435000,
        "customers": 856,
    }
    return render_template(
        "dashboard.html",
        stock_items=stock,
        alerts=alerts,
        stats=stats,
        kpis=kpis,
        posts=NETWORK_POSTS,
        transfers=TRANSFERS,
        banner_text=banner_text,
    )


@app.get("/app/inventory")
def inventory_page():
    # Inventory page uses uploaded stock data when present; otherwise it falls back to the seeded mock data.
    return render_template("inventory.html", stock_items=with_status(get_active_stock()))


@app.get("/app/customers")
def customers_page():
    # Customer list page renders the current mock CRM records.
    return render_template("customers.html", customers=CUSTOMERS)


@app.get("/app/sales")
def sales_page():
    # Sales page renders uploaded sales data when present; otherwise it falls back to the seeded mock dataset.
    uploaded_sales = session.get("uploaded_sales")
    return render_template("sales.html", sales_orders=uploaded_sales if uploaded_sales else SALES_ORDERS)


@app.get("/app/analytics")
def analytics_page():
    # Analytics page spreads the stock and transfer KPIs through the template.
    stats = compute_stats(STOCK_LEVELS, TRANSFERS)
    uploaded_sales = session.get("uploaded_sales")
    if uploaded_sales:
        revenue = sum(float(order.get("amount", 0) or 0) for order in uploaded_sales)
        order_count = len(uploaded_sales)
        avg_order_value = revenue / order_count if order_count else 0
        kpis = {
            "revenue": revenue,
            "profit": 435000,
            "orders": order_count,
            "avg_order_value": avg_order_value,
        }
    else:
        kpis = {
            "revenue": 1875000,
            "profit": 435000,
            "orders": 1246,
            "avg_order_value": 1500,
        }
    return render_template("analytics.html", stats=stats, kpis=kpis)


@app.get("/app/team")
def team_page():
    # Team roster uses the mock employee dataset.
    return render_template("team.html", team=TEAM)


@app.get("/app/network")
def network_page():
    # Network feed is powered by the mock social posts.
    return render_template("network.html", posts=NETWORK_POSTS)


@app.get("/app/settings")
def settings_page():
    # Settings page is static but still served by the Flask app.
    return render_template("settings.html")


if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000)
